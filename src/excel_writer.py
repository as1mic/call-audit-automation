import openpyxl

from copy import copy
from datetime import datetime
from pathlib import Path
from openpyxl.styles import Font


DATE_COLUMN = 1
REQUEST_TYPE_COLUMN = 2
PHONE_COLUMN = 3
MANAGER_COLUMN = 5
INTRO_COLUMN = 6
BODY_COLUMN = 7
YEAR_COLUMN = 8
MILEAGE_COLUMN = 9
DIAGNOSTIC_OFFER_COLUMN = 10
PREVIOUS_REPAIRS_COLUMN = 11
BOOKING_DATE_COLUMN = 12
GOODBYE_COLUMN = 13
TOP_JOB_COLUMN = 14
TOP_JOB_OK_COLUMN = 15
RECOMMENDATION_COLUMN = 16
RESULT_COLUMN = 17
SCORE_COLUMN = 18
PARTS_COLUMN = 19
COMMENT_COLUMN = 20
TOTAL_SCORE_COLUMN = 21
TEMPLATE_ROW = 3



def load_report(report_path: Path):
    return openpyxl.load_workbook(report_path)


def save_report(workbook, destination_path: Path) -> Path:
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination_path)
    return destination_path


def get_main_sheet(workbook):
    return workbook.worksheets[0]


def normalize_phone(phone: str) -> str:
    digits = []

    for char in str(phone):
        if char.isdigit():
            digits.append(char)

    phone_digits = "".join(digits)

    if len(phone_digits) >= 10:
        return phone_digits[-10:]

    return phone_digits


def normalize_date(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    date_text = str(value).strip()

    for date_format in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(date_text, date_format).strftime("%Y-%m-%d")
        except ValueError:
            pass

    return date_text


def copy_cell_style(source_cell, target_cell):
    if source_cell.has_style:
        target_cell._style = copy(source_cell._style)

    target_cell.font = copy(source_cell.font)
    target_cell.fill = copy(source_cell.fill)
    target_cell.border = copy(source_cell.border)
    target_cell.alignment = copy(source_cell.alignment)
    target_cell.protection = copy(source_cell.protection)
    target_cell.number_format = source_cell.number_format


def copy_row_style(sheet, source_row: int, target_row: int):
    sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height

    for column in range(1, max(sheet.max_column, TOTAL_SCORE_COLUMN) + 1):
        copy_cell_style(sheet.cell(row=source_row, column=column), sheet.cell(row=target_row, column=column))


def prepare_report_sheet(sheet):
    header_cell = sheet.cell(row=2, column=TOTAL_SCORE_COLUMN)

    if not header_cell.value:
        copy_cell_style(sheet.cell(row=2, column=SCORE_COLUMN), header_cell)
        header_cell.value = "Балів"

    if not sheet.cell(row=1, column=TOTAL_SCORE_COLUMN).value:
        copy_cell_style(sheet.cell(row=1, column=COMMENT_COLUMN), sheet.cell(row=1, column=TOTAL_SCORE_COLUMN))

    sheet.column_dimensions[openpyxl.utils.get_column_letter(TOTAL_SCORE_COLUMN)].width = 12


def find_row_by_date_and_phone(sheet, target_date, target_phone) -> int:
    target_date = str(target_date)
    target_phone = normalize_phone(target_phone)

    for row in range(3, sheet.max_row + 1):
        date_value = normalize_date(sheet.cell(row=row, column=DATE_COLUMN).value)
        phone_value = normalize_phone(sheet.cell(row=row, column=PHONE_COLUMN).value)

        if date_value == target_date and phone_value == target_phone:
            return row

    return -1


def get_next_empty_row(sheet) -> int:
    for row in range(3, sheet.max_row + 1):
        date_value = sheet.cell(row=row, column=DATE_COLUMN).value
        phone_value = sheet.cell(row=row, column=PHONE_COLUMN).value

        if date_value is None and phone_value is None:
            return row

    return sheet.max_row + 1


def get_or_create_row(sheet, metadata: dict) -> int:
    row_number = find_row_by_date_and_phone(sheet, metadata["date"], metadata["phone"])

    if row_number != -1:
        return row_number

    row_number = get_next_empty_row(sheet)
    copy_row_style(sheet, TEMPLATE_ROW, row_number)

    return row_number


def write_text_phone(sheet, row_number: int, phone: str):
    phone_cell = sheet.cell(row=row_number, column=PHONE_COLUMN)
    phone_cell.value = normalize_phone(phone)
    phone_cell.number_format = "@"


def write_analysis_to_row(sheet, row_number, metadata: dict, analysis_result):
    flags = analysis_result.get("flags", [])
    has_booking = analysis_result.get("has_booking", False)
    manager_score = analysis_result.get("manager_score", 0)

    sheet.cell(row=row_number, column=DATE_COLUMN).value = metadata.get("date", "")
    sheet.cell(row=row_number, column=REQUEST_TYPE_COLUMN).value = analysis_result.get("request_type", "")
    write_text_phone(sheet, row_number, metadata.get("phone", ""))

    sheet.cell(row=row_number, column=MANAGER_COLUMN).value = analysis_result.get("manager_name", "")
    sheet.cell(row=row_number, column=INTRO_COLUMN).value = analysis_result.get("greeting_score", 0)
    sheet.cell(row=row_number, column=BODY_COLUMN).value = analysis_result.get("body_score", 0)
    sheet.cell(row=row_number, column=YEAR_COLUMN).value = analysis_result.get("year_score", 0)
    sheet.cell(row=row_number, column=MILEAGE_COLUMN).value = analysis_result.get("mileage_score", 0)
    sheet.cell(row=row_number, column=DIAGNOSTIC_OFFER_COLUMN).value = analysis_result.get("diagnostic_offer_score", 0)
    sheet.cell(row=row_number, column=PREVIOUS_REPAIRS_COLUMN).value = analysis_result.get("previous_repairs_score", 0)
    sheet.cell(row=row_number, column=BOOKING_DATE_COLUMN).value = ""
    sheet.cell(row=row_number, column=GOODBYE_COLUMN).value = analysis_result.get("goodbye_score", 0)
    sheet.cell(row=row_number, column=TOP_JOB_COLUMN).value = analysis_result.get("top_job", "")
    sheet.cell(row=row_number, column=TOP_JOB_OK_COLUMN).value = "Так" if analysis_result.get("top_job_ok", 0) else "Ні"
    sheet.cell(row=row_number, column=RECOMMENDATION_COLUMN).value = ", ".join(flags)
    sheet.cell(row=row_number, column=RESULT_COLUMN).value = "Є запис" if has_booking else "Немає запису"
    sheet.cell(row=row_number, column=SCORE_COLUMN).value = manager_score
    sheet.cell(row=row_number, column=PARTS_COLUMN).value = ""
    sheet.cell(row=row_number, column=TOTAL_SCORE_COLUMN).value = manager_score

    comment_cell = sheet.cell(row=row_number, column=COMMENT_COLUMN)
    comment_cell.value = analysis_result.get("comment", "")

    if flags:
        comment_cell.font = Font(color="FF0000")
    else:
        comment_cell.font = Font(color="000000")
